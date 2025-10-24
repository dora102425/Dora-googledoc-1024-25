import os
import io
import time
import json
import zipfile
import re
import base64
from typing import List, Dict, Any, Optional

# External Libraries - Make sure these are in requirements.txt
import streamlit as st
import pandas as pd
import yaml
import mammoth
from docx import Document
from PIL import Image

# Try to import provider SDKs, handle gracefully if not installed
try:
    import google.generativeai as genai
except ImportError:
    genai = None
try:
    from openai import OpenAI
except ImportError:
    OpenAI = None
try:
    from xai_sdk import Client
    from xai_sdk.chat import user as grok_user, system as grok_system
except ImportError:
    Client, grok_user, grok_system = None, None, None


# ==============================================================================
# --- SERVICES (CORRECTED with Bug Fix) ---
# ==============================================================================

class ProviderError(Exception):
    pass

# --- Gemini Service ---
def run_gemini(model: str, system_prompt: str, user_prompt: str, input_text: str, temperature: float, max_tokens: int, images: Optional[List[Any]] = None) -> str:
    api_key = st.session_state.get("google_api_key")
    if not api_key: raise ProviderError("Google API Key is not set. Please provide it in the sidebar.")
    if not genai: raise ProviderError("google-generativeai is not installed.")
    
    genai.configure(api_key=api_key)
    user_prompt_filled = user_prompt.replace("{{input}}", input_text)
    full_prompt = f"{system_prompt}\n\n{user_prompt_filled}"
    
    parts = [full_prompt]
    if images:
        for f in images:
            f.seek(0)
            parts.append(Image.open(io.BytesIO(f.read())))

    m = genai.GenerativeModel(model)
    resp = m.generate_content(parts, generation_config={"temperature": temperature, "max_output_tokens": max_tokens})
    return resp.text

# --- OpenAI Service ---
def run_openai(model: str, system_prompt: str, user_prompt: str, input_text: str, temperature: float, max_tokens: int, images: Optional[List[Any]] = None) -> str:
    api_key = st.session_state.get("openai_api_key")
    if not api_key: raise ProviderError("OpenAI API Key is not set. Please provide it in the sidebar.")
    if not OpenAI: raise ProviderError("openai is not installed.")

    client = OpenAI(api_key=api_key)
    content_items = [{"type": "text", "text": user_prompt.replace("{{input}}", input_text)}]
    if images:
        for f in images:
            f.seek(0)
            b64 = base64.b64encode(f.read()).decode("utf-8")
            content_items.append({"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}})

    msgs = [{"role": "system", "content": system_prompt}] if system_prompt else []
    msgs.append({"role": "user", "content": content_items})
    
    completion = client.chat.completions.create(model=model, messages=msgs, temperature=temperature, max_tokens=max_tokens)
    return completion.choices[0].message.content or ""

# --- Grok Service ---
def run_grok(model: str, system_prompt: str, user_prompt: str, input_text: str, temperature: float, max_tokens: int, images: Optional[List[Any]] = None) -> str:
    api_key = st.session_state.get("xai_api_key")
    if not api_key: raise ProviderError("XAI (Grok) API Key is not set. Please provide it in the sidebar.")
    if not Client: raise ProviderError("xai-sdk is not installed.")
    if images: st.warning("Grok provider does not support images in this app. Ignoring.")

    client = Client(api_key=api_key)
    chat = client.chat.create(model=model)
    if system_prompt: chat.append(grok_system(system_prompt))
    chat.append(grok_user(user_prompt.replace("{{input}}", input_text)))
    
    response = chat.sample(max_len=max_tokens, temp=temperature)
    return getattr(response, "content", "") or ""

# --- Model Router (THE FIX IS HERE) ---
def run_agent_step(agent_config: dict, input_text: str, images: Optional[List[Any]]) -> str:
    """
    Filters arguments from the agent config and calls the correct provider function.
    """
    provider = agent_config.get("provider")

    # Prepare a clean dictionary of arguments for the provider functions
    provider_args = {
        "model": agent_config.get("model"),
        "system_prompt": agent_config.get("system_prompt", ""),
        "user_prompt": agent_config.get("user_prompt", "{{input}}"),
        "input_text": input_text,
        "temperature": float(agent_config.get("temperature", 0.5)),
        "max_tokens": int(agent_config.get("max_tokens", 2048)),
        "images": images
    }

    try:
        if provider == "gemini":
            return run_gemini(**provider_args)
        elif provider == "openai":
            return run_openai(**provider_args)
        elif provider == "grok":
            return run_grok(**provider_args)
        else:
            raise ProviderError(f"Unknown provider: {provider}")
    except Exception as e:
        # Catch and re-raise as a standard ProviderError for consistent handling
        raise ProviderError(str(e))


# ==============================================================================
# --- UTILS (Parser, Template, Exporter) ---
# ==============================================================================
def parse_dataset_file(file) -> List[Dict[str, Any]]:
    name = file.name.lower()
    if name.endswith(".csv"): df = pd.read_csv(file)
    elif name.endswith(".json"): return json.load(file)
    elif name.endswith(".xlsx"): df = pd.read_excel(file)
    elif name.endswith(".ods"): df = pd.read_excel(file, engine="odf")
    elif name.endswith(".txt"): return [{"text": l} for l in file.getvalue().decode("utf-8").splitlines() if l.strip()]
    else: raise ValueError("Unsupported file type")
    return df.fillna("").to_dict(orient="records")

def parse_template_file(file) -> str:
    name = file.name.lower()
    if name.endswith((".txt", ".md")): return file.getvalue().decode("utf-8")
    elif name.endswith(".docx"): return mammoth.extract_raw_text(file).value
    raise ValueError("Unsupported template type")

def guess_schema(records: List[Dict[str, Any]]) -> List[str]:
    cols = set()
    for r in records[:50]: cols.update(r.keys())
    return sorted(list(cols))

def render_template(template: str, record: Dict[str, Any]) -> str:
    return re.sub(r"\{\{([^}]+)\}\}", lambda m: str(record.get(m.group(1).strip(), m.group(0))), template)

def export_docx(text: str) -> io.BytesIO:
    doc, buf = Document(), io.BytesIO()
    for line in str(text).splitlines(): doc.add_paragraph(line)
    doc.save(buf)
    buf.seek(0)
    return buf

def build_zip(docs: List[Dict[str,str]], ext: str) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for i, d in enumerate(docs):
            fname = os.path.splitext(d.get("filename") or f"doc_{i+1}")[0] + f".{ext}"
            content = d.get("content", "")
            if ext == "docx": zf.writestr(fname, export_docx(content).getvalue())
            else: zf.writestr(fname, content)
    buf.seek(0)
    return buf.getvalue()


# ==============================================================================
# --- STREAMLIT APP ---
# ==============================================================================
st.set_page_config(page_title="Agentic Docs Builder", page_icon="🌸", layout="wide")
ss = st.session_state

# --- UI & THEMES ---
THEMES = { "Flora (Default)": ("#A55EEA", "#F0F2F6", "#E6E6FA", "#262730"), "Ocean Breeze": ("#1E90FF", "#F0F8FF", "#D6EAF8", "#17202A"), "Forest Whisper": ("#228B22", "#F0FFF0", "#D5F5E3", "#145A32"), "Sunset Glow": ("#FF4500", "#FFF5EE", "#FADBD8", "#641E16"), "Midnight Slate": ("#34495E", "#2C3E50", "#5D6D7E", "#ECF0F1"), "Sunny Meadow": ("#FFD700", "#FFFFF0", "#FCF3CF", "#873600"), "Lavender Dreams": ("#8A2BE2", "#F8F5FF", "#E8DAEF", "#300D4F"), "Ruby Red": ("#C70039", "#FFF0F0", "#F5B7B1", "#581845"), "Emerald Isle": ("#009B77", "#E8F8F5", "#A9DFBF", "#0E6655"), "Golden Harvest": ("#DAA520", "#FFFAF0", "#FDEBD0", "#7E5109"), "Arctic Ice": ("#5DADE2", "#EBF5FB", "#D4E6F1", "#154360"), "Warm Earth": ("#A0522D", "#FFF8DC", "#F5DEB3", "#512E0C"), "Cyberpunk Neon": ("#00FFFF", "#1B2631", "#283747", "#EAECEE"), "Pastel Cloud": ("#FFB6C1", "#FFF9FA", "#FADADD", "#6C3483"), "Volcanic Ash": ("#36454F", "#212F3D", "#2C3E50", "#FDFEFE"), "Mint Fresh": ("#66CDAA", "#F0FFFA", "#D1F2EB", "#0B5345"), "Berry Fusion": ("#8B0000", "#FAEBD7", "#F5CBA7", "#4A235A"), "Grape Vine": ("#6A0DAD", "#F4ECF7", "#D7BDE2", "#2C1B4B"), "Teal Focus": ("#008080", "#E0FFFF", "#B2DFDB", "#004D40"), "Mustard Zing": ("#FFDB58", "#FFFACD", "#F9E79F", "#5C4033")}
def get_theme_css(theme_name):
    primary, bg, sec_bg, text = THEMES.get(theme_name, THEMES["Flora (Default)"])
    return f"""<style> .badge {{padding:2px 8px;border-radius:8px;font-size:12px;display:inline-block;margin-right:6px;}} .badge-ok {{background:#e6ffe6;color:#1a7f37;border:1px solid #1a7f37;}} .badge-err {{background:#ffe6e6;color:#a12622;border:1px solid #a12622;}} .status-dot {{height:10px;width:10px;border-radius:50%;display:inline-block;margin-right:6px;}} .dot-green {{background:#00c853;}} .dot-yellow {{background:#ffd600;}} .dot-red {{background:#d50000;}} .panel {{padding:10px 12px;border:1px solid #ccc;border-radius:8px;background:{sec_bg};}} .stApp {{background-color:{bg};}} h1,h2,h3,p,label,div[data-baseweb="tooltip"],div[data-testid="stMarkdownContainer"] p {{color:{text} !important;}} .stButton>button {{background-color:{primary};color:white !important;border:1px solid {primary};}} div[data-testid="stSidebarUserContent"] {{background-color:{sec_bg};}}</style>"""

# --- State Init ---
DEFAULT_AGENTS = yaml.safe_load("""- {name: Summarizer, provider: gemini, model: gemini-1.5-flash, temperature: 0.3, max_tokens: 1024, system_prompt: 'Summarize accurately.', user_prompt: 'Summarize in 3 bullets:\\n\\n{{input}}'}""")
if "agents" not in ss: ss.agents = DEFAULT_AGENTS
for key in ["dataset", "schema", "generated_docs", "pipeline_history", "images"]:
    if key not in ss: ss[key] = []
for key, val in {"template_text": "", "is_running": False, "selected_theme": "Flora (Default)"}.items():
    if key not in ss: ss[key] = val
st.markdown(get_theme_css(ss.selected_theme), unsafe_allow_html=True)

# --- SIDEBAR ---
with st.sidebar:
    st.header("🌸 Flora Controls")
    st.selectbox("UI Theme", list(THEMES.keys()), key="selected_theme")
    st.divider()

    st.caption("API Keys")
    def api_key_manager(env_var, ss_key, label):
        if env_var in os.environ:
            ss[ss_key] = os.environ[env_var]
        is_set = ss_key in ss and ss[ss_key]
        badge = f"<span class='badge {'badge-ok' if is_set else 'badge-err'}'>{label}</span>"
        if not is_set:
            key_input = st.text_input(f"Enter {label} API Key", type="password", key=f"in_{ss_key}")
            if key_input:
                ss[ss_key] = key_input
                st.rerun()
        return badge
    c1, c2, c3 = st.columns(3)
    c1.markdown(api_key_manager("GOOGLE_API_KEY", "google_api_key", "Gemini"), unsafe_allow_html=True)
    c2.markdown(api_key_manager("OPENAI_API_KEY", "openai_api_key", "OpenAI"), unsafe_allow_html=True)
    c3.markdown(api_key_manager("XAI_API_KEY", "xai_api_key", "Grok"), unsafe_allow_html=True)
    st.divider()

    st.caption("Global Overrides")
    ALL_MODELS = ["None", "gemini-1.5-flash", "gemini-1.5-pro", "gpt-4o", "gpt-4o-mini", "llama3-70b-8192", "llama3-8b-8192"]
    prov = st.selectbox("Provider override", ["None", "gemini", "openai", "grok"])
    mod = st.selectbox("Model override", ALL_MODELS)
    ss.global_provider_override = prov if prov != "None" else None
    ss.global_model_override = mod if mod != "None" else None
    st.divider()

    st.caption("Agent Configuration")
    yaml_file = st.file_uploader("Load agents.yaml", type=["yaml", "yml"])
    if yaml_file:
        try:
            agents = yaml.safe_load(yaml_file)
            if isinstance(agents, list): ss.agents = agents
            else: st.error("YAML must be a list of agents.")
        except Exception as e: st.error(f"YAML Error: {e}")
    st.download_button("Download agents.yaml", yaml.safe_dump(ss.agents), "agents.yaml")
    st.divider()
    
    ss.images = st.file_uploader("Images for Vision Models", type=["png","jpg","jpeg","webp"], accept_multiple_files=True)

# --- MAIN UI ---
st.title("Agentic Docs Builder – Flora Edition")
c1,c2,c3,c4 = st.columns(4); c1.metric("Records", len(ss.dataset)); c2.metric("Schema Cols", len(ss.schema)); c3.metric("Docs", len(ss.generated_docs)); c4.metric("Agents", len(ss.agents))
tab1, tab2, tab3, tab4, tab5, tab6 = st.tabs(["1. Data", "2. Template", "3. Docs", "4. Agents", "5. Pipeline", "6. Export"])

with tab1:
    ds_file = st.file_uploader("Upload Dataset (CSV, JSON, XLSX, ODS, TXT)", type=["csv","json","xlsx","ods","txt"])
    if ds_file:
        try:
            ss.dataset = parse_dataset_file(ds_file); ss.schema = guess_schema(ss.dataset)
            st.success(f"Loaded {len(ss.dataset)} records."); st.dataframe(pd.DataFrame(ss.dataset).head(), use_container_width=True)
        except Exception as e: st.error(f"Parse Error: {e}")

with tab2:
    c1, c2 = st.columns(2)
    with c1:
        tpl_file = st.file_uploader("Upload Template (.txt, .md, .docx)", type=["txt","md","docx"])
        if tpl_file: ss.template_text = parse_template_file(tpl_file)
        st.text_area("Template", key="template_text", height=300, help="Use {{column_name}} placeholders")
        if st.button("Generate Docs", type="primary", disabled=(not ss.dataset or not ss.template_text)):
            ss.generated_docs = [ {"filename": r.get("filename",f"doc_{i}.txt"), "content": render_template(ss.template_text, r)} for i, r in enumerate(ss.dataset, 1)]
            st.success(f"Generated {len(ss.generated_docs)} docs.")
    with c2:
        st.markdown("#### Live Preview")
        if ss.dataset and ss.template_text: st.markdown(f"<div class='panel'>{render_template(ss.template_text, ss.dataset[0])[:2000]}</div>", unsafe_allow_html=True)

with tab3:
    if not ss.generated_docs: st.info("No documents generated yet.")
    for i, doc in enumerate(ss.generated_docs):
        with st.expander(f"Doc {i+1}: {doc.get('filename')}"):
            ss.generated_docs[i]["filename"] = st.text_input("Filename", doc.get("filename"), key=f"fn_{i}")
            ss.generated_docs[i]["content"] = st.text_area("Content", doc.get("content"), height=250, key=f"fc_{i}")

with tab4:
    for i, agent in enumerate(ss.agents):
        with st.expander(f"Agent {i+1} • {agent.get('name', 'Agent')}"):
            c1,c2,c3 = st.columns(3)
            agent["name"] = c1.text_input("Name", agent.get("name"), key=f"an_{i}")
            agent["provider"] = c2.selectbox("Provider", ["gemini","openai","grok"], index=["gemini","openai","grok"].index(agent.get("provider","gemini")), key=f"ap_{i}")
            agent["model"] = c3.selectbox("Model", ALL_MODELS, index=ALL_MODELS.index(agent.get("model")) if agent.get("model") in ALL_MODELS else 0, key=f"am_{i}")
            agent["temperature"] = st.slider("Temp", 0.0, 2.0, agent.get("temperature",0.5), 0.05, key=f"at_{i}")
            agent["max_tokens"] = st.number_input("Max Tokens", 128, 16384, agent.get("max_tokens",2048), key=f"amt_{i}")
            agent["system_prompt"] = st.text_area("System Prompt", agent.get("system_prompt"), height=100, key=f"asp_{i}")
            agent["user_prompt"] = st.text_area("User Prompt", agent.get("user_prompt"), height=150, key=f"aup_{i}")
    c1, c2 = st.columns(2)
    if c1.button("➕ Add Agent"): ss.agents.append(DEFAULT_AGENTS[0]); st.rerun()
    if c2.button("🔄 Reset"): ss.agents = DEFAULT_AGENTS; st.rerun()

with tab5: # Pipeline execution with the bug fix
    input_text = st.text_area("Pipeline Input", ss.generated_docs[0]['content'] if ss.generated_docs else "", height=250)
    if st.button("Run Pipeline", type="primary", disabled=(ss.is_running or not ss.agents)):
        ss.is_running, ss.pipeline_history = True, []
        progress, status = st.progress(0), st.empty()
        current_text = input_text
        for i, agent in enumerate(ss.agents, 1):
            agent_copy = agent.copy()
            if ss.global_provider_override: agent_copy["provider"] = ss.global_provider_override
            if ss.global_model_override: agent_copy["model"] = ss.global_model_override
            status.markdown(f"<span class='status-dot dot-yellow'></span> Running {i}/{len(ss.agents)}: **{agent_copy['name']}**", unsafe_allow_html=True)
            progress.progress(i/len(ss.agents))
            try:
                # This is the corrected function call
                output = run_agent_step(agent_config=agent_copy, input_text=current_text, images=ss.images)
                ss.pipeline_history.append({"agent": agent_copy['name'], "input": current_text, "output": output, "error": None})
                current_text = output
            except ProviderError as e:
                ss.pipeline_history.append({"agent": agent_copy['name'], "input": current_text, "output": None, "error": str(e)})
                st.error(f"Pipeline Error: {e}"); break
        status.markdown("<span class='status-dot dot-green'></span> Pipeline finished.", unsafe_allow_html=True)
        ss.is_running = False
    if ss.pipeline_history:
        st.markdown("--- \n### Pipeline Results")
        for i, step in enumerate(ss.pipeline_history):
            with st.expander(f"Step {i+1}: <{step['agent']}>", expanded=bool(step["error"])):
                st.text_area("Input", step["input"], height=150, disabled=True, key=f"in_{i}")
                if step["error"]: st.error(step["error"])
                else: st.text_area("Output", step["output"], height=150, disabled=True, key=f"out_{i}")

with tab6:
    if not ss.generated_docs: st.info("Generate docs to export.")
    else:
        fmt = st.selectbox("Format", ["txt", "md", "docx", "zip-txt", "zip-md", "zip-docx"])
        if "zip" in fmt:
            if st.button("Build ZIP", type="primary"):
                ss.last_export = build_zip(ss.generated_docs, ext=fmt.split("-")[1])
                st.success("ZIP ready.")
            if "last_export" in ss: st.download_button("Download ZIP", ss.last_export, f"docs_{int(time.time())}.zip")
        else:
            idx = st.number_input("Doc index", 1, len(ss.generated_docs), 1) - 1
            doc = ss.generated_docs[idx]
            fname = os.path.splitext(doc["filename"])[0] + f".{fmt}"
            content = export_docx(doc["content"]) if fmt == "docx" else doc["content"].encode("utf-8")
            st.download_button(f"Download .{fmt.upper()}", content, fname)
